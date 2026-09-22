from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
OUTPUT = ROOT / "Framework de Validacao de Textura CT Pulmonar.docx"


def shade(cell, color):
    cell_properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), color)
    cell_properties.append(shading)


def cell_margins(cell, top=100, start=110, bottom=100, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_border(cell, color="D9D9D9"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:color"), color)


def prevent_row_split(row):
    properties = row._tr.get_or_add_trPr()
    no_split = OxmlElement("w:cantSplit")
    properties.append(no_split)


def repeat_header_row(row):
    properties = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    properties.append(header)


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = False
    header_cells = table.rows[0].cells
    repeat_header_row(table.rows[0])
    for i, label in enumerate(headers):
        header_cells[i].text = label
        shade(header_cells[i], "17365D")
        if widths:
            header_cells[i].width = Inches(widths[i])
        for run in header_cells[i].paragraphs[0].runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.size = Pt(9)
        header_cells[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        cell_margins(header_cells[i])
        set_cell_border(header_cells[i])
    for row_index, row_values in enumerate(rows):
        row = table.add_row()
        prevent_row_split(row)
        cells = row.cells
        for i, value in enumerate(row_values):
            cells[i].text = value
            if widths:
                cells[i].width = Inches(widths[i])
            if row_index % 2 == 1:
                shade(cells[i], "F3F6FA")
            for paragraph in cells[i].paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.size = Pt(9)
            cells[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            cell_margins(cells[i])
            set_cell_border(cells[i])
    doc.add_paragraph().paragraph_format.space_after = Pt(3)
    return table


def add_bullet(doc, text):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.28)
    paragraph.paragraph_format.first_line_indent = Inches(-0.22)
    paragraph.paragraph_format.space_before = Pt(1)
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.add_run("• ")
    paragraph.add_run(text)
    return paragraph


def add_numbered(doc, number, text):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.28)
    paragraph.paragraph_format.first_line_indent = Inches(-0.24)
    paragraph.paragraph_format.space_before = Pt(1)
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.add_run(f"{number}. ").bold = True
    paragraph.add_run(text)
    return paragraph


def main():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.12
    for name, size in (("Title", 22), ("Heading 1", 15), ("Heading 2", 12)):
        style = styles[name]
        style.font.name = "Aptos Display" if name == "Title" else "Aptos"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), style.font.name)
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(11 if name != "Title" else 0)
        style.paragraph_format.space_after = Pt(5)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.add_run("Framework de Validacao de Textura em CT Pulmonar Sintetica")
    subtitle = doc.add_paragraph()
    subtitle.add_run("Documento de contexto para o projeto ARIA").bold = True
    subtitle.runs[0].font.size = Pt(12)
    subtitle.runs[0].font.color.rgb = RGBColor(89, 89, 89)
    doc.add_paragraph(
        "Este documento registra a primeira definição do framework que decidirá se uma CT pulmonar sintética pode ser usada em data augmentation. "
        "A decisão combina fidelidade ao caso original que condicionou a geração e compatibilidade com a variabilidade observada em imagens reais validadas do mesmo dataset."
    )

    doc.add_heading("Objetivo", level=1)
    doc.add_paragraph(
        "O framework avalia uma imagem sintética individualmente, sempre junto da imagem original que a condicionou. "
        "A imagem deve preservar a textura do caso, continuar plausível para a população real de referência e não introduzir sinais de geração artificial, como alisamento excessivo, repetição de padrões ou perda de detalhes finos."
    )
    doc.add_paragraph(
        "Nesta versão inicial, o método é calibrado para o dataset atual. Ele não é, ainda, uma regra universal para qualquer CT ou protocolo de aquisição."
    )

    doc.add_heading("Regiao analisada e pre processamento", level=1)
    doc.add_paragraph("A análise é mask based: a textura é medida somente no parênquima pulmonar.")
    for item in (
        "A máscara pulmonar da imagem original define a região de interesse.",
        "Uma pequena erosão remove a borda pulmão tecido externo, uma transição intensa que poderia dominar as métricas.",
        "A mesma máscara espacial é aplicada à sintética pareada, pois a geração foi condicionada pela original e as duas estão alinhadas.",
        "Clipping de intensidade, discretização, direções e distância entre pixels permanecem fixos para toda comparação.",
        "No cálculo de vizinhança, só entram pares de pixels que estejam dentro da máscara."
    ):
        add_bullet(doc, item)

    doc.add_heading("Entradas e referencia interna", level=1)
    add_table(doc, ["Papel", "Conteudo"], [
        ("Entrada", "Uma imagem sintética e a imagem original que a condicionou."),
        ("Referencia interna", "Perfil estatístico construído previamente a partir de imagens originais validadas."),
        ("Saida", "Métricas, comparação com o par, comparação com a referência e decisão de uso.")
    ], widths=[1.55, 5.9])
    doc.add_paragraph(
        "O perfil de referência fica embutido e versionado no framework. Ele contém as distribuições das métricas nas imagens reais aceitas e os parâmetros usados no pré-processamento. "
        "No piloto atual, a referência utiliza 37 imagens originais centrais, com os dois pulmões e máscaras revisadas operacionalmente."
    )

    doc.add_heading("Metricas de textura", level=1)
    doc.add_paragraph("O conjunto inicial reúne quatro famílias complementares. Cada uma observa um aspecto diferente da textura.")

    doc.add_heading("GLCM relacao entre pixels vizinhos", level=2)
    doc.add_paragraph("A GLCM descreve como tons de cinza de pixels próximos aparecem juntos. Ela captura a microtextura local.")
    add_bullet(doc, "Contraste: mede a variação entre vizinhos. Valor menor na sintética pode indicar textura lisa demais e perda de heterogeneidade; valor maior pode indicar ruído ou artefato.")
    add_bullet(doc, "Energia: mede o predomínio de padrões locais repetidos. Valor alto demais pode indicar uma textura regular ou estampada; valor muito baixo pode representar perda de organização local.")

    doc.add_heading("GLRLM continuidade de padroes", level=2)
    doc.add_paragraph("A GLRLM mede sequências de pixels vizinhos com intensidade parecida. Ela acrescenta a noção de continuidade dos padrões.")
    add_bullet(doc, "Ênfase em runs curtos: candidata para representar detalhes finos e pequenas variações.")
    add_bullet(doc, "Ênfase em runs longos: candidata para identificar áreas amplas e uniformes; valor alto demais pode reforçar o sinal de alisamento.")

    doc.add_heading("Power Spectrum distribuicao por escala", level=2)
    doc.add_paragraph("O espectro de potência mostra se a energia visual está mais concentrada em estruturas amplas ou em detalhes finos.")
    add_bullet(doc, "Inclinação espectral: resume o equilíbrio entre componentes de grande escala e altas frequências.")
    add_bullet(doc, "Fração de energia em alta frequência: queda excessiva sugere perda de detalhe; aumento excessivo pode ser granulação ou ruído artificial.")

    doc.add_heading("Wavelet textura localizada em varias escalas", level=2)
    doc.add_paragraph("Wavelet separa a imagem em componentes de baixa e alta frequência em diferentes escalas, mantendo a localização dos detalhes.")
    add_bullet(doc, "Proporção de energia nos componentes de alta frequência: candidata para medir preservação de detalhe local.")
    add_bullet(doc, "Entropia dos componentes: candidata para medir diversidade e complexidade dos padrões de detalhe.")
    doc.add_paragraph("Wavelet é uma transformação, e não uma métrica única. As features finais serão escolhidas depois da etapa de avaliação.")

    doc.add_heading("Como a decisao funciona", level=1)
    doc.add_heading("Eixo A fidelidade ao caso original", level=2)
    doc.add_paragraph(
        "Para cada feature, comparamos a sintética à sua original: diferença do par igual a feature da sintética menos feature da original. "
        "A pergunta é: a geração preservou a textura específica daquele pulmão?"
    )
    doc.add_heading("Eixo B compatibilidade com imagens reais", level=2)
    doc.add_paragraph(
        "A feature da sintética é posicionada na distribuição das originais validadas por meio de mediana, dispersão, percentil e faixas de alerta. "
        "A pergunta é: a imagem permanece plausível dentro da população real deste dataset?"
    )
    doc.add_paragraph(
        "O objetivo não é ficar igual à média. Um pulmão real pode ser raro e ainda válido; por isso, a posição da original na referência também contextualiza a sintética."
    )

    doc.add_heading("Casos possiveis e decisao para data augmentation", level=1)
    add_table(doc,
        ["Par", "Referencia real", "Leitura", "Retorno", "Uso"],
        [
            ("Próxima", "Dentro da faixa", "Preservou o caso e continua realista.", "Aprovada", "Entra"),
            ("Próxima", "Original e sintética raras de modo compatível", "Caso incomum, mas preservado.", "Aprovada com observação", "Entra com etiqueta"),
            ("Distante", "Dentro da faixa", "Pode ter ocorrido regressão à média ou mudança relevante.", "Revisão manual", "Fica fora até revisão"),
            ("Distante", "Fora da faixa", "Provável falha de geração ou artefato.", "Rejeitada", "Não entra")
        ], widths=[0.75, 1.35, 2.1, 1.45, 1.5]
    )
    doc.add_paragraph(
        "A imagem original só pode condicionar uma sintética se já tiver sido validada. Uma original com ruído ou artefato não deve legitimar uma sintética semelhante."
    )

    doc.add_heading("Saida esperada por imagem", level=1)
    for item in (
        "Valores de cada feature na original e na sintética.",
        "Diferença do par e interpretação.",
        "Percentil da original e da sintética na referência real.",
        "Sinalização por feature: normal, atenção ou fora da faixa.",
        "Decisão final: aprovada, aprovada com observação, revisão manual ou rejeitada.",
        "Motivo legível da decisão, por exemplo: contraste menor que o par e abaixo da referência, sugerindo alisamento excessivo."
    ):
        add_bullet(doc, item)

    doc.add_heading("Proximos passos avaliacao das features", level=1)
    doc.add_paragraph("Antes de uma feature entrar definitivamente no framework, ela deverá atender aos quatro critérios abaixo.")
    for index, item in enumerate((
        "Complementaridade: medir algo que as features já escolhidas não capturam bem.",
        "Estabilidade: reagir de forma consistente a pequenas mudanças razoáveis de máscara, erosão, discretização e janela de intensidade.",
        "Baixa redundância: não ser quase uma cópia de outra feature. Isso será verificado pela correlação entre features nas originais validadas.",
        "Interpretação clara: uma alteração precisa permitir uma hipótese visual útil, como alisamento, repetição, perda de detalhe, continuidade excessiva ou ruído."
    ), start=1):
        add_numbered(doc, index, item)
    doc.add_paragraph(
        "Também verificaremos se cada feature varia o suficiente entre imagens reais válidas e se consegue identificar transformações controladas, como suavização ou adição de ruído, sem reagir aleatoriamente."
    )

    doc.add_heading("Ordem de implementacao", level=1)
    for index, item in enumerate((
        "Consolidar a referência GLCM nas originais validadas com contraste e energia.",
        "Definir a normalização da diferença entre sintética e original para cada feature.",
        "Implementar o relatório por par para GLCM, inicialmente sem uma decisão automática definitiva.",
        "Avaliar candidatas de GLRLM, Power Spectrum e Wavelet pelos quatro critérios.",
        "Incorporar apenas as features aprovadas e definir a regra integrada de aprovação, revisão e rejeição.",
        "Gerar exemplos visuais de cada caso para a apresentação futura."
    ), start=1):
        add_numbered(doc, index, item)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.add_run("Projeto ARIA  |  Framework de validação de textura")
    footer_run.font.size = Pt(8)
    footer_run.font.color.rgb = RGBColor(100, 100, 100)

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
